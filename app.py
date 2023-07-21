import tkinter as tk
from tkinter import messagebox
from docx import Document

def generate_word_document(responses):
    doc = Document("template.docx")  # Assuming you have a template file named "template.docx" in the same directory.

    # Replace the placeholders in the template with the user's responses
    for i, response in enumerate(responses, start=1):
        placeholder = f"{{answer_{i}}}"
        for paragraph in doc.paragraphs:
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, response)

    doc.save("output.docx")
    messagebox.showinfo("Success", "The Word document has been generated successfully!")

def submit_form():
    responses = [entry.get() for entry in entry_fields]
    generate_word_document(responses)

# Create the GUI
root = tk.Tk()
root.title("Form and Word Document Generator")

# Create input fields for the form
entry_labels = ["Question 1", "Question 2", "Question 3"]  # Replace with your questions.
entry_fields = []
for i, label_text in enumerate(entry_labels):
    label = tk.Label(root, text=label_text)
    label.pack()
    entry = tk.Entry(root)
    entry.pack()
    entry_fields.append(entry)

# Submit button
submit_button = tk.Button(root, text="Submit", command=submit_form)
submit_button.pack()

root.mainloop()
